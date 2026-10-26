
from pathlib import Path as _WorkspacePath
import os as _workspace_os
RESEARCH_WORKSPACE = _WorkspacePath(_workspace_os.environ.get("RESEARCH_WORKSPACE", _workspace_os.getcwd())).resolve()

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document(str(RESEARCH_WORKSPACE / 'input-manuscript.docx'))

# Extract text with structure
output = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        style = para.style.name if para.style else "Normal"
        output.append(f"[{i}] [{style}] {text}")

# Also extract tables
table_text = []
for ti, table in enumerate(doc.tables):
    table_text.append(f"\n=== Table {ti+1} ===")
    for row in table.rows:
        row_text = [cell.text.strip() for cell in row.cells]
        table_text.append(" | ".join(row_text))

with open(str(RESEARCH_WORKSPACE / 'paper_content.txt'), "w", encoding="utf-8") as f:
    f.write("=== PARAGRAPHS ===\n")
    f.write("\n".join(output))
    f.write("\n\n=== TABLES ===\n")
    f.write("\n".join(table_text))

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print("Saved to paper_content.txt")
